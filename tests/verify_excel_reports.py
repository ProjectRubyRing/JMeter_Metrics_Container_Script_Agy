import glob
import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

files = sorted(glob.glob("tests/sample_reports/*.xlsx"))
print(f"Found {len(files)} Excel reports:")
assert len(files) == 4, f"Expected 4 xlsx files, got {len(files)}"

expected_tokens = [
    ("フロントコンテナ", "intra-webサービス"),
    ("バックコンテナ", "intra-apiサービス"),
    ("フロントコンテナ", "inter-apiサービス"),
    ("バックコンテナ", "sf-apiサービス"),
]

expected_sheets = [
    "サマリー",
    "1_DB接続プール_JBossCLI",
    "2_ヒープ_GC_JFR",
    "3_top_ps_プロセス状況",
    "4_ActiveProcessorCount要否",
    "5_詳細ログ_jcmd_JFR_top",
]

for f in files:
    bname = os.path.basename(f)
    print(f"  [FILE] {bname}")
    wb = openpyxl.load_workbook(f)
    print(f"    Sheets: {wb.sheetnames}")
    assert wb.sheetnames == expected_sheets, f"Unexpected sheet names: {wb.sheetnames}"
    for name in wb.sheetnames:
        ws = wb[name]
        a1 = ws["A1"].value
        print(f"      - {name}: {ws.max_row} rows x {ws.max_column} cols | A1={a1!r}")
    # Check data row in sheet 1..4
    for s_idx in range(1, 5):
        ws = wb[expected_sheets[s_idx]]
        assert ws.max_row >= 7, f"Expected data rows in {expected_sheets[s_idx]}, got max_row={ws.max_row}"
    wb.close()

for role_jp, svc_jp in expected_tokens:
    matched = [f for f in files if role_jp in os.path.basename(f) and svc_jp in os.path.basename(f)]
    assert len(matched) == 1, f"Missing expected report filename for {role_jp} + {svc_jp}"

print("=== ALL OPENPYXL EXCEL REPORT VERIFICATIONS PASSED! ===")
